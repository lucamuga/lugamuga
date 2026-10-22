"""Build an Excel journal for the 10:00 ET reversal: enter a few prices per day from the 5m chart,
the workbook computes direction, entry, stop, target, P&L and a monthly overview.

    python backtest/make_excel_template.py --from 2026-09-22 --to 2026-10-21 --out 10-Uhr-Reversal.xlsx
"""

from __future__ import annotations

import argparse
from datetime import date

import pandas as pd
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

FONT = "Arial"
FIRST, LAST = 5, 204  # data rows on the Journal sheet
INPUT_FILL = PatternFill("solid", start_color="FFF2CC")
HEAD_FILL = PatternFill("solid", start_color="1F3864")
CALC_FILL = PatternFill("solid", start_color="F2F2F2")
THIN = Border(bottom=Side(style="thin", color="D9D9D9"))
BLUE = Font(name=FONT, color="0000FF")
BLACK = Font(name=FONT)
BOLD = Font(name=FONT, bold=True)
HEAD = Font(name=FONT, bold=True, color="FFFFFF")
TITLE = Font(name=FONT, bold=True, size=14)

PRICE, PCT, USD, R = "#,##0.00", "0.00%", '#,##0 "$";-#,##0 "$";"-"', '0.00"R";-0.00"R";"-"'

SETTINGS = [
    # name, label, value, format, note
    ("MinDrive", "Mindestbewegung 09:30–10:00", 0.002, PCT, "Kleinere Bewegungen werden nicht gehandelt."),
    ("ZielR", "Ziel in R (Vielfaches des Risikos)", 2, "0.0", "2 = Ziel liegt doppelt so weit weg wie der Stop."),
    ("Punktwert", "$ pro Punkt und Kontrakt", 20, USD, "NQ = 20, MNQ = 2, QQQ-Aktie = 1 (dann Kontrakte = Stückzahl)."),
    ("Kontrakte", "Kontrakte / Stück", 1, "0", ""),
    ("Kommission", "Gebühr pro Kontrakt (Hin + Rück) in $", 4, USD, "Je nach Broker; MNQ etwa 1,50 $."),
    ("Slippage", "Slippage pro Seite in Punkten", 0.25, PRICE, "1 Tick NQ = 0,25 Punkte; wird bei Ein- und Ausstieg abgezogen."),
]

# (header, width, kind, number format, comment)
COLUMNS = [
    ("Datum", 11, "in", "DD.MM.YYYY", None),
    ("Open 09:30", 11, "in", PRICE, "Eröffnungskurs der 09:30-Kerze"),
    ("Close 09:55-Kerze", 11, "in", PRICE, "Schlusskurs der letzten Kerze vor 10:00"),
    ("Tageshoch bis Einstieg", 12, "in", PRICE, "Höchster Kurs seit 09:30 bis vor die Einstiegskerze (= Stop bei Short)"),
    ("Tagestief bis Einstieg", 12, "in", PRICE, "Tiefster Kurs seit 09:30 bis vor die Einstiegskerze (= Stop bei Long)"),
    ("Hoch 10:00-Kerze", 11, "in", PRICE, "Hoch der 5-Minuten-Kerze von 10:00"),
    ("Tief 10:00-Kerze", 11, "in", PRICE, "Tief der 5-Minuten-Kerze von 10:00"),
    ("Einstiegskurs", 11, "in", PRICE,
     "Schließt die 10:05- oder 10:10-Kerze gegen die Morgenrichtung außerhalb der 10:00-Kerze: "
     "Open der nächsten Kerze eintragen. Kein Ausbruch bis 10:15 = leer lassen."),
    ("Ergebnis", 9, "in", None, "Was wurde zuerst erreicht? Ziel, Stop oder keins von beiden bis 11:30 (= Zeit)."),
    ("Kurs 11:30", 10, "in", PRICE, "Nur bei Ergebnis = Zeit: Schlusskurs der 11:25-Kerze"),
    ("Notiz", 18, "in", None, "z. B. News um 10:00 (ISM, JOLTS ...), Fed-Rede"),
    ("Morgen-Bewegung", 10, "calc", PCT, None),
    ("Richtung", 10, "calc", None, None),
    ("Status", 10, "calc", None, None),
    ("Stop", 11, "calc", PRICE, None),
    ("Risiko (Pkt)", 9, "calc", PRICE, None),
    ("Ziel", 11, "calc", PRICE, None),
    ("Exitkurs", 11, "calc", PRICE, None),
    ("Punkte netto", 9, "calc", PRICE, None),
    ("R", 8, "calc", R, None),
    ("Gewinn/Verlust", 11, "calc", USD, None),
    ("Kumuliert", 11, "calc", USD, None),
    ("Höchststand", 11, "calc", USD, None),
    ("Drawdown", 11, "calc", USD, None),
    ("Hinweis", 16, "calc", None, None),
    ("Wochentag", 6, "calc", "0", None),
]


def journal_formulas(r: int) -> list[str]:
    d = 'IF(M{r}="Long",1,-1)'.format(r=r)
    return [
        f'=IF(OR(B{r}="",C{r}=""),"",(C{r}-B{r})/B{r})',
        f'=IF(L{r}="","",IF(ABS(L{r})<MinDrive,"kein Trade",IF(L{r}>0,"Short","Long")))',
        f'=IF(OR(M{r}="",M{r}="kein Trade"),"",IF(H{r}="","kein Signal","Trade"))',
        f'=IF(N{r}<>"Trade","",IF(M{r}="Short",D{r},E{r}))',
        f'=IF(N{r}<>"Trade","",ABS(H{r}-O{r}))',
        f'=IF(N{r}<>"Trade","",H{r}+{d}*ZielR*P{r})',
        f'=IF(OR(N{r}<>"Trade",I{r}=""),"",IF(I{r}="Ziel",Q{r},IF(I{r}="Stop",O{r},IF(J{r}="","",J{r}))))',
        f'=IF(R{r}="","",(R{r}-H{r})*{d}-2*Slippage)',
        f'=IF(OR(S{r}="",P{r}=0),"",S{r}/P{r})',
        f'=IF(S{r}="","",S{r}*Punktwert*Kontrakte-Kommission*Kontrakte)',
        f'=IF(U{r}="","",SUM($U${FIRST}:U{r}))',
        f'=IF(V{r}="","",MAX(0,MAX($V${FIRST}:V{r})))',
        f'=IF(V{r}="","",V{r}-W{r})',
        f'=IF(N{r}<>"Trade","",IF(P{r}=0,"Stop = Einstieg?",IF(I{r}="","Ergebnis eintragen",'
        f'IF(AND(I{r}="Zeit",J{r}=""),"Kurs 11:30 fehlt",IF(OR(AND(M{r}="Short",H{r}>F{r}),'
        f'AND(M{r}="Long",H{r}<G{r})),"Einstieg prüfen","ok")))))',
        f'=IF(A{r}="","",WEEKDAY(A{r},2))',
    ]


def rng(col: str) -> str:
    return f"Journal!${col}${FIRST}:${col}${LAST}"


def build(days: list[date], path: str) -> None:
    wb = Workbook()

    # --- Anleitung ---------------------------------------------------------------
    ws = wb.active
    ws.title = "Anleitung"
    ws.column_dimensions["A"].width = 4
    ws.column_dimensions["B"].width = 110
    lines = [
        ("Nasdaq 10:00-Reversal: Trading-Journal (5-Minuten-Chart)", TITLE),
        ("", None),
        ("Die Strategie", BOLD),
        ("1. Morgen-Bewegung: Wie weit läuft der Nasdaq (NQ/MNQ oder QQQ) von 09:30 bis 10:00 New Yorker Zeit "
         "(in Deutschland meist 15:30 bis 16:00)?", None),
        ("2. Ist die Bewegung groß genug (Einstellungen: Mindestbewegung), wird um 10:00 dagegen gehandelt: "
         "nach einem Anstieg Short, nach einem Fall Long.", None),
        ("3. Einstieg: Schließt die 10:05- oder 10:10-Kerze unter dem Tief (bei Short) bzw. über dem Hoch (bei Long) "
         "der 10:00-Kerze, wird zum Open der nächsten Kerze eingestiegen.", None),
        ("4. Stop: Tageshoch (Short) bzw. Tagestief (Long) bis zum Einstieg. Ziel: 2 × Risiko. "
         "Spätestens um 11:30 wird geschlossen.", None),
        ("", None),
        ("So füllst du das Journal aus (Blatt „Journal“)", BOLD),
        ("• Nur die gelben Zellen mit blauer Schrift ausfüllen. Graue Spalten rechnet Excel selbst.", None),
        ("• Pro Handelstag eine Zeile. Die Daten des letzten Monats sind schon eingetragen; "
         "weitere Tage einfach unten anhängen.", None),
        ("• Kurse aus dem 5-Minuten-Chart ablesen (TradingView, Broker). Alle Kurse vom selben Instrument nehmen.", None),
        ("• Gibt es keinen Ausbruch bis 10:15, bleibt „Einstiegskurs“ leer. Der Tag zählt dann als „kein Signal“.", None),
        ("• Nach dem Eintragen des Einstiegs zeigt Excel Stop und Ziel. Im Chart nachsehen, was zuerst erreicht wurde, "
         "und bei „Ergebnis“ Ziel, Stop oder Zeit wählen. Bei Zeit den Kurs um 11:30 eintragen.", None),
        ("• Die Spalte „Hinweis“ meldet fehlende Angaben und Einstiege, die nicht zur Regel passen.", None),
        ("• Kontraktgröße, Gebühren und Regeln lassen sich im Blatt „Einstellungen“ ändern. Die Übersicht passt sich an.", None),
        ("", None),
        ("Beispiel (erfundene Zahlen, nur zur Veranschaulichung)", BOLD),
        ("Open 09:30 = 20.000 · Close 09:55 = 20.090 (+0,45 % → Short) · Tageshoch bis Einstieg = 20.110 · "
         "Tagestief = 19.995 · 10:00-Kerze Hoch 20.110 / Tief 20.070", None),
        ("Die 10:05-Kerze schließt bei 20.064, also unter 20.070 → Einstieg zum Open der 10:10-Kerze: 20.062. "
         "Excel zeigt: Stop 20.110, Risiko 48 Punkte, Ziel 19.966.", None),
        ("Der Kurs fällt bis 10:40 auf 19.966 → Ergebnis „Ziel“ → ca. +95,5 Punkte netto → mit NQ ca. +1.906 $ (+1,99 R).", None),
        ("", None),
        ("Wichtig", BOLD),
        ("Kein Anlageratschlag. Ein Monat (ca. 20 Handelstage) zeigt, wie es zuletzt lief, beweist aber keinen "
         "dauerhaften Vorteil. Für eine belastbare Aussage braucht es mehrere Jahre Daten (siehe backtest/README.md).", None),
    ]
    for i, (text, font) in enumerate(lines, start=1):
        c = ws.cell(row=i, column=2, value=text)
        c.font = font or BLACK
        c.alignment = Alignment(wrap_text=True, vertical="top")

    # --- Einstellungen -----------------------------------------------------------
    st = wb.create_sheet("Einstellungen")
    st["A1"] = "Einstellungen"
    st["A1"].font = TITLE
    st["A2"] = "Gelbe Zellen mit blauer Schrift dürfen geändert werden."
    st["A2"].font = Font(name=FONT, italic=True, color="808080")
    for col, width in zip("ABC", (38, 12, 70)):
        st.column_dimensions[col].width = width
    for i, (name, label, value, fmt, note) in enumerate(SETTINGS, start=4):
        st.cell(row=i, column=1, value=label).font = BLACK
        c = st.cell(row=i, column=2, value=value)
        c.font, c.fill, c.number_format = BLUE, INPUT_FILL, fmt
        st.cell(row=i, column=3, value=note).font = Font(name=FONT, color="595959")
        wb.defined_names[name] = DefinedName(name, attr_text=f"Einstellungen!$B${i}")
    st.cell(row=4 + len(SETTINGS) + 1, column=1,
            value="Standardwerte entsprechen dem Python-Backtest (backtest/strategy.py): 0,2 % Mindestbewegung, 2R-Ziel, "
                  "Exit 11:30.").font = Font(name=FONT, italic=True, color="808080")

    # --- Journal -----------------------------------------------------------------
    jn = wb.create_sheet("Journal")
    jn["A1"] = "Journal: eine Zeile pro Handelstag"
    jn["A1"].font = TITLE
    jn["A2"] = "Gelb = selbst eintragen (aus dem 5-Minuten-Chart). Grau = rechnet Excel. Zeiten = New York (ET)."
    jn["A2"].font = Font(name=FONT, italic=True, color="808080")
    jn["A3"], jn["L3"] = "EINGABE", "BERECHNUNG"
    jn["A3"].font = jn["L3"].font = BOLD
    for j, (header, width, kind, fmt, note) in enumerate(COLUMNS, start=1):
        letter = jn.cell(row=4, column=j).column_letter
        h = jn.cell(row=4, column=j, value=header)
        h.font, h.fill = HEAD, HEAD_FILL
        h.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        if note:
            h.comment = Comment(note, "Journal")
        jn.column_dimensions[letter].width = width
    jn.row_dimensions[4].height = 32
    jn.column_dimensions["Z"].hidden = True

    for r in range(FIRST, LAST + 1):
        k = r - FIRST
        for j, (_, _, kind, fmt, _) in enumerate(COLUMNS, start=1):
            c = jn.cell(row=r, column=j)
            if kind == "in":
                c.font, c.fill = BLUE, INPUT_FILL
            else:
                c.font, c.fill = BLACK, CALC_FILL
            if fmt:
                c.number_format = fmt
            c.border = THIN
        if k < len(days):
            jn.cell(row=r, column=1, value=days[k])
        for j, formula in enumerate(journal_formulas(r), start=12):
            jn.cell(row=r, column=j, value=formula)

    dv = DataValidation(type="list", formula1='"Ziel,Stop,Zeit"', allow_blank=True)
    dv.error, dv.errorTitle = "Bitte Ziel, Stop oder Zeit wählen.", "Ergebnis"
    jn.add_data_validation(dv)
    dv.add(f"I{FIRST}:I{LAST}")
    green, red = Font(name=FONT, color="00703C"), Font(name=FONT, color="C00000")
    for col in ("U", "T"):
        jn.conditional_formatting.add(f"{col}{FIRST}:{col}{LAST}", CellIsRule(operator="greaterThan", formula=["0"], font=green))
        jn.conditional_formatting.add(f"{col}{FIRST}:{col}{LAST}", CellIsRule(operator="lessThan", formula=["0"], font=red))
    jn.freeze_panes = f"B{FIRST}"

    # --- Übersicht ---------------------------------------------------------------
    ov = wb.create_sheet("Übersicht", 1)
    ov["A1"] = "Übersicht"
    ov["A1"].font = TITLE
    ov["A2"] = "Alles automatisch aus dem Journal berechnet."
    ov["A2"].font = Font(name=FONT, italic=True, color="808080")
    for col, width in zip("ABCDEF", (30, 14, 14, 14, 14, 14)):
        ov.column_dimensions[col].width = width

    A, N, I, U, T, M, Z = (rng(c) for c in "ANIUTMZ")
    stats = [
        ("Zeitraum von", f'=IF(COUNT({A})=0,"",MIN({A}))', "DD.MM.YYYY"),
        ("Zeitraum bis", f'=IF(COUNT({A})=0,"",MAX({A}))', "DD.MM.YYYY"),
        ("Erfasste Tage", f"=COUNT(Journal!$L${FIRST}:$L${LAST})", "0"),
        ("Tage mit Trade", f'=COUNTIF({N},"Trade")', "0"),
        ("davon abgeschlossen", f"=COUNT({U})", "0"),
        ("Gewinner", f'=COUNTIF({U},">0")', "0"),
        ("Verlierer", f'=COUNTIF({U},"<=0")', "0"),
        ("Trefferquote", '=IF(B8=0,"",B9/B8)', "0.0%"),
        ("Gewinn/Verlust gesamt", f"=SUM({U})", USD),
        ("Ø pro Trade", f'=IF(B8=0,"",AVERAGE({U}))', USD),
        ("Ø R pro Trade", f'=IF(COUNT({T})=0,"",AVERAGE({T}))', R),
        ("Profit Factor", f'=IF(SUMIF({U},"<0")=0,"",SUMIF({U},">0")/-SUMIF({U},"<0"))', "0.00"),
        ("Größter Gewinn", f'=IF(B8=0,"",MAX({U}))', USD),
        ("Größter Verlust", f'=IF(B8=0,"",MIN({U}))', USD),
        ("Max. Drawdown", f"=MIN(Journal!$X${FIRST}:$X${LAST})", USD),
        ("Exit am Ziel", f'=COUNTIFS({N},"Trade",{I},"Ziel")', "0"),
        ("Exit am Stop", f'=COUNTIFS({N},"Trade",{I},"Stop")', "0"),
        ("Exit um 11:30", f'=COUNTIFS({N},"Trade",{I},"Zeit")', "0"),
    ]
    for i, (label, formula, fmt) in enumerate(stats, start=4):
        ov.cell(row=i, column=1, value=label).font = BLACK
        c = ov.cell(row=i, column=2, value=formula)
        c.font, c.number_format = BOLD, fmt
    assert ov["A8"].value == "davon abgeschlossen" and ov["A9"].value == "Gewinner"

    def table(top: int, title: str, keys: list, key_label: str, crit) -> int:
        ov.cell(row=top, column=1, value=title).font = Font(name=FONT, bold=True, size=12)
        for j, h in enumerate([key_label, "Trades", "Gewinner", "Trefferquote", "Gewinn/Verlust", "Ø R"], start=1):
            c = ov.cell(row=top + 1, column=j, value=h)
            c.font, c.fill = HEAD, HEAD_FILL
        for k, key in enumerate(keys):
            r = top + 2 + k
            label_cell, cond = crit(r, key)
            ov.cell(row=r, column=1, value=label_cell[0]).number_format = label_cell[1]
            guard = f'IF($A{r}="","",' if label_cell[2] else "("
            ov.cell(row=r, column=2, value=f"={guard}COUNTIFS({U},\">-1E+99\"{cond}))")
            ov.cell(row=r, column=3, value=f"={guard}COUNTIFS({U},\">0\"{cond}))")
            ov.cell(row=r, column=4, value=f'=IF(OR($A{r}="",N(B{r})=0),"",C{r}/B{r})').number_format = "0.0%"
            ov.cell(row=r, column=5, value=f"={guard}SUMIFS({U}{cond}))").number_format = USD
            ov.cell(row=r, column=6, value=f'=IF(OR($A{r}="",N(B{r})=0),"",'
                                            f'SUMIFS({T}{cond})/B{r})').number_format = R
            for j in range(1, 7):
                ov.cell(row=r, column=j).font = BLACK
        return top + 2 + len(keys) + 1

    def month_crit(r, k):
        if k == 0:
            label = f'=IF(COUNT({A})=0,"",DATE(YEAR(MIN({A})),MONTH(MIN({A})),1))'
        else:
            label = f'=IF(A{r - 1}="","",IF(EDATE(A{r - 1},1)>MAX({A}),"",EDATE(A{r - 1},1)))'
        return (label, "MMMM YYYY", True), f',{A},">="&$A{r},{A},"<"&EDATE($A{r},1)'

    nxt = table(24, "Nach Monat", list(range(12)), "Monat", month_crit)
    weekdays = ["Montag", "Dienstag", "Mittwoch", "Donnerstag", "Freitag"]
    nxt = table(nxt, "Nach Wochentag", weekdays, "Wochentag",
                lambda r, k: ((k, "@", False), f",{Z},{weekdays.index(k) + 1}"))
    table(nxt, "Nach Richtung", ["Long", "Short"], "Richtung",
          lambda r, k: ((k, "@", False), f',{M},"{k}"'))

    wb.move_sheet("Journal", offset=-2)  # order: Anleitung, Journal, Übersicht, Einstellungen
    wb.save(path)


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--from", dest="start", type=date.fromisoformat, required=True)
    ap.add_argument("--to", dest="end", type=date.fromisoformat, required=True)
    ap.add_argument("--out", default="10-Uhr-Reversal-Journal.xlsx")
    a = ap.parse_args()
    days = [d.date() for d in pd.bdate_range(a.start, a.end)]
    build(days, a.out)
    print(f"Wrote {a.out} with {len(days)} prefilled trading days")


if __name__ == "__main__":
    main()
